import requests
from pathlib import Path
import pdfplumber
from docx import Document
import globals

def extract_text_from_pdf(file_path):
    text = []
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text.append(page_text)
    return "\n".join(text)

def extract_text_from_docx(file_path):
    doc = Document(file_path)
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())

def summarize_resume(resume_file):
    resume_text = ""
    if resume_file.endswith(".pdf"):
        resume_text = extract_text_from_pdf(resume_file)
    else:
        resume_text = extract_text_from_docx(resume_file)
    return generate_response(globals.summarize_resume_prompt_filepath, additional_prompt_text=resume_text)

def load_prompt(prompt_filepath):
    file_path = Path(prompt_filepath)
    if not file_path.exists():
        print(f"Prompt file does not exist: {prompt_filepath}")
        return ""

    with file_path.open(encoding="utf-8") as f:
        prompt = f.read()

    # Replace placeholders
    prompt = prompt.replace("{job}", globals.job_choice)
    prompt = prompt.replace("{nb_question}", str(globals.max_question_amount))
    prompt = prompt.replace("{resume_summary}", globals.resume_summary)
    prompt = prompt.replace(
        "{chat_history}",
        "\n".join(
            f"{'User' if m['role']=='user' else 'Assistant'}: {m['content']}"
            for m in globals.chat_history
        )
    )
    return prompt

def generate_response(prompt_filepath, additional_prompt_text=""):
    prompt = load_prompt(prompt_filepath) + additional_prompt_text
    print(prompt)
    try:
        response = requests.post(globals.generative_ai_url, json={
            "model": globals.generative_ai_model,
            "prompt": prompt,
            "stream": False
        })
        response.raise_for_status()
        return response.json()["response"]
    except Exception as e:
        print(f"Erreur lors de la génération de la réponse : {e}")
        return "Erreur lors de la génération de la réponse."

